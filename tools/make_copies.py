# -*- coding: utf-8 -*-
"""Build A4 PDF + Word copies of the FIN211 cheat sheet from the HTML source."""
import io, re, os
from html.parser import HTMLParser

SRC = os.path.join('..', 'fin211-cheatsheet.html')
OUT_PDF = os.path.join('..', 'site', 'FIN211_Cheat_Sheet_A4.pdf')
OUT_DOCX = os.path.join('..', 'site', 'FIN211_Cheat_Sheet.docx')

cs = io.open(SRC, encoding='utf-8').read()
# expand collapsibles + drop nav for print/docx
cs = cs.replace('<details>', '<div class="prntx">').replace('</details>', '</div>')
cs = cs.replace('<summary>', '<p class="sumh">').replace('</summary>', '</p>')
cs = re.sub(r'<nav>.*?</nav>', '', cs, flags=re.S)

# ---------------- PDF (WeasyPrint, A4) ----------------
extra = """
@page { size: A4; margin: 13mm 11mm; }
section{box-shadow:none!important;border:1px solid #dde4ee!important;page-break-inside:auto!important;margin-bottom:10px!important}
h2,h3{page-break-after:avoid}
.prntx{page-break-inside:avoid;margin:6px 0}
.sumh{font-weight:700;color:#8a6d1a;margin:6px 0 2px}
body{font-size:11.5px}
.f{font-size:12.5px}
table{font-size:10.5px}
tr,.f,.callout,.prntx,.term,.card,.flow,.rule{page-break-inside:avoid!important}
table{page-break-inside:auto}
h2,h3,h4,.sumh{page-break-after:avoid}
p,li{orphans:3;widows:3}
.cols{columns:1!important}
section{page-break-inside:auto!important;page-break-before:always!important}
section:first-of-type{page-break-before:auto!important}
header.hero{page-break-after:avoid}
"""
pdf_html = cs.replace('</style>', extra + '</style>')
from weasyprint import HTML
HTML(string=pdf_html).write_pdf(OUT_PDF)
print('PDF written:', OUT_PDF)

# ---------------- DOCX (python-docx) ----------------
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SUP = {'2': '²', '3': '³', '4': '⁴', '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹', '0': '⁰', '1': '¹'}
def supmap(t): return ''.join(SUP.get(ch, ch) for ch in t)
def wrapfrac(side):
    side = side.strip()
    return '(' + side + ')' if re.search(r'[ ±×÷,;]', side) else side

class Conv(HTMLParser):
    def __init__(self, doc):
        super().__init__(convert_charrefs=True)
        self.doc = doc
        self.runs = []          # list of [text, bold, italic]
        self.bold = 0; self.ital = 0
        self.skip = 0
        self.para_kind = None   # 'p','li','h1','h2','h3','sumh','f'
        self.fr = []            # fraction stack: dicts {top:..., bot:..., mode:}
        self.table = None; self.row = None; self.cell = None
        self.buf = lambda: self.cell if self.cell is not None else self.runs
    def addtext(self, t):
        if not t: return
        if self.skip: return
        if self.fr:
            key = self.fr[-1]['mode']
            if key: self.fr[-1][key] += t
            return
        self.buf().append([t, self.bold > 0, self.ital > 0])
    def flush(self, kind):
        runs = self.runs; self.runs = []
        text = ''.join(r[0] for r in runs).strip()
        if not text and not runs: return
        if kind == 'h1':
            return  # hero title already added manually
        if kind == 'h2':
            hp = self.doc.add_heading(text, 1)
            if not getattr(self, '_first_h2', False):
                self._first_h2 = True
            else:
                hp.paragraph_format.page_break_before = True
            hp.paragraph_format.keep_with_next = True
            return
        if kind == 'h3':
            hp = self.doc.add_heading(text, 2)
            hp.paragraph_format.keep_with_next = True
            return
        p = self.doc.add_paragraph(style='List Bullet' if kind == 'li' else None)
        if kind == 'f':
            p.paragraph_format.left_indent = Mm(6)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
        if not runs:
            p.add_run(text)
            return
        for t, b, i in runs:
            r = p.add_run(t)
            r.bold = b; r.italic = i
            if kind == 'f':
                r.font.name = 'Cambria'
                r.font.size = Pt(11)
        if kind == 'sumh':
            for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor(0x8a, 0x6d, 0x1a)
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        cls = a.get('class', '')
        if tag in ('b', 'strong'): self.bold += 1
        elif tag in ('i', 'em'): self.ital += 1
        elif tag == 'a' and 'top-link' in cls: self.skip += 1
        elif tag == 'sub': pass
        elif tag == 'sup': self.addtext('^')
        elif tag == 'span' and 'fr' in cls: self.fr.append({'top': '', 'bot': '', 'mode': None})
        elif tag == 'span' and 'top' in cls and self.fr: self.fr[-1]['mode'] = 'top'
        elif tag == 'span' and 'bot' in cls and self.fr: self.fr[-1]['mode'] = 'bot'
        elif tag == 'p': self.para_kind = 'sumh' if 'sumh' in cls else 'p'
        elif tag == 'li': self.para_kind = 'li'
        elif tag in ('h1', 'h2', 'h3'): self.para_kind = tag
        elif tag == 'div' and 'f' in cls: self.para_kind = 'f'
        elif tag == 'table': self.table = []
        elif tag == 'tr': self.row = []
        elif tag in ('td', 'th'): self.cell = []
    def handle_endtag(self, tag):
        if tag in ('b', 'strong'): self.bold = max(0, self.bold - 1)
        elif tag in ('i', 'em'): self.ital = max(0, self.ital - 1)
        elif tag == 'a': self.skip = max(0, self.skip - 1)
        elif tag == 'span' and self.fr and self.fr[-1]['mode'] in ('top', 'bot'): self.fr[-1]['mode'] = None
        elif tag == 'span' and self.fr and self.fr[-1]['mode'] is None and (self.fr[-1]['top'] or self.fr[-1]['bot']):
            f = self.fr.pop()
            self.addtext(' ' + wrapfrac(f['top']) + ' / ' + wrapfrac(f['bot']) + ' ')
        elif tag == 'p': self.flush(self.para_kind or 'p'); self.para_kind = None
        elif tag == 'li': self.flush('li'); self.para_kind = None
        elif tag in ('h1', 'h2', 'h3'): self.flush(tag); self.para_kind = None
        elif tag == 'div' and self.para_kind == 'f': self.flush('f'); self.para_kind = None
        elif tag in ('td', 'th'):
            txt = ''.join(r[0] for r in self.cell).strip()
            self.row.append(txt); self.cell = None
        elif tag == 'tr':
            if self.row is not None: self.table.append(self.row)
            self.row = None
        elif tag == 'table':
            if self.table:
                ncols = max(len(r) for r in self.table)
                t = self.doc.add_table(rows=0, cols=ncols)
                t.style = 'Table Grid'
                for ri, r in enumerate(self.table):
                    new_row = t.add_row()
                    new_row.cant_split = True
                    cells = new_row.cells
                    for ci in range(ncols):
                        cells[ci].text = r[ci] if ci < len(r) else ''
                        if ri == 0:
                            for pp in cells[ci].paragraphs:
                                for rr in pp.runs: rr.bold = True
                self.doc.add_paragraph('')
            self.table = None
    def handle_data(self, data):
        t = data.replace('\xa0', ' ')
        if self.para_kind in ('h1', 'h2', 'h3', 'p', 'li', 'f', 'sumh') or self.cell is not None or self.fr:
            self.addtext(t)
        elif t.strip() and self.table is None and self.para_kind is None:
            pass

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(15)
sec.top_margin = sec.bottom_margin = Mm(14)
doc.add_heading('FIN211 Corporate Finance — Exam Cheat Sheet', 0)
sub = doc.add_paragraph('All formulas, decision rules, key definitions and worked examples from the course lectures (Lectures 1–13).')
sub.runs[0].italic = True
Conv(doc).feed(cs)
doc.save(OUT_DOCX)
print('DOCX written:', OUT_DOCX)
