# -*- coding: utf-8 -*-
"""Build printable A4 PDF + Word: ALL formulas, abbreviations and definitions (bilingual)."""
import io, os, re, json
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = json.load(io.open(os.path.join(HERE, 'data.json'), encoding='utf-8'))
OUT_PDF = os.path.join(HERE, '..', 'site', 'FIN211_Formulas_Abbrev_Definitions_A4.pdf')
OUT_DOCX = os.path.join(HERE, '..', 'site', 'FIN211_Formulas_Abbrev_Definitions.docx')

CH = {c['id']: c for c in DATA['chapters']}
FI, AI, DI = {}, {}, {}
LS = {l['id']: l for l in DATA['lessons']}

SUPSUP = {'2': '²', '3': '³', '4': '⁴', '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹', '0': '⁰', '1': '¹', 'n': 'ⁿ'}
def _strip(t): return re.sub(r'<[^>]+>', '', t)
def lin(html):
    """Linearize formula HTML to plain text for Word."""
    s = html
    s = re.sub(r'<sub>(.*?)</sub>', lambda m: _strip(m.group(1)), s, flags=re.S)
    s = re.sub(r'<sup>(.*?)</sup>', lambda m: SUPSUP.get(_strip(m.group(1)), '^' + _strip(m.group(1))), s, flags=re.S)
    def fr(m):
        top, bot = _strip(m.group(1)).strip(), _strip(m.group(2)).strip()
        wt = '(' + top + ')' if re.search(r'[ ±×÷,;]', top) else top
        wb = '(' + bot + ')' if re.search(r'[ ±×÷,;]', bot) else bot
        return wt + ' / ' + wb
    s = re.sub(r'<span class="fr"><span class="top">(.*?)</span><span class="bot">(.*?)</span></span>', fr, s, flags=re.S)
    s = s.replace('<br>', '  ·  ')
    s = re.sub(r'<[^>]+>', '', s)
    s = (s.replace('&gt;', '>').replace('&lt;', '<').replace('&amp;', '&')
          .replace('&nbsp;', ' ').replace('&#8722;', '−'))
    return re.sub(r'\s+', ' ', s).strip()

SUPPL = [  # definitions present in lessons but not yet in the glossary
 ("NFAI", "Net Fixed Asset Investment", "Change in net fixed assets plus depreciation — the cash absorbed by the fixed-asset base.", "صافي الاستثمار في الأصول الثابتة", "التغير في صافي الأصول الثابتة زائد الإهلاك — النقد الذي تمتصه قاعدة الأصول الثابتة.", "l4"),
 ("NCAI", "Net Current Asset Investment", "Change in current assets less the change in payables and accruals.", "صافي الاستثمار في الأصول المتداولة", "التغير في الأصول المتداولة ناقص التغير في الموردين والمصروفات المستحقة.", "l4"),
 ("Cash-flow margin", "Cash-flow margin", "OCF ÷ sales revenue — the cash generated from each revenue dollar.", "هامش التدفق النقدي", "OCF ÷ إيرادات المبيعات — النقد المولّد من كل دولار إيرادات.", "l5"),
 ("Cash-flow coverage", "Cash-flow coverage", "OCF ÷ total debt — the ability to repay debt from operating cash.", "تغطية التدفق النقدي", "OCF ÷ إجمالي الدين — القدرة على سداد الدين من نقد التشغيل.", "l5"),
 ("Dividend yield", "Dividend yield", "D ÷ P — the dividend component of a stock's expected return.", "عائد التوزيع", "D ÷ P — مكوّن التوزيعات في العائد المتوقع للسهم.", "l10"),
 ("Capital gain", "Capital gain", "The appreciation of the stock's market price due to earnings growth.", "الربح الرأسمالي", "ارتفاع سعر السوق للسهم بفعل نمو الأرباح.", "l10"),
 ("Required return (rs)", "Required return", "The return investors demand to hold an asset; the discount rate of valuation models.", "العائد المطلوب", "العائد الذي يطلبه المستثمرون لحمل الأصل؛ وهو معدل الخصم في نماذج التقييم.", "l10"),
 ("Correlation coefficient (ρ)", "Correlation coefficient", "How two assets move together, from +1 (perfect unison) to −1 (perfect opposite).", "معامل الارتباط", "كيف يتحرك أصلان معًا، من +1 (تناغم تام) إلى −1 (تعاكس تام).", "l12"),
 ("Residual claim", "Residual claim", "Common shareholders are paid only after every other claim is satisfied.", "المطالبة المتبقية", "لا يُدفع للمساهمين العاديين إلا بعد استيفاء كل المطالب الأخرى.", "l9"),
 ("Levered / unlevered OCF", "Levered vs unlevered OCF", "After interest and debt servicing (net income + Dep) vs before it (EBIT(1−T) + Dep).", "التدفق التشغيلي الممول / غير الممول", "بعد الفوائد وخدمة الدين (صافي الدخل + الإهلاك) مقابل قبلها (EBIT(1−T) + الإهلاك).", "l5"),
]

# ============================== HTML for PDF ==============================
CSS = """
@page { size: A4; margin: 13mm 11mm 15mm;
  @bottom-center { content: counter(page) " / " counter(pages); font-size:9px; color:#5c6b84; font-family:Tahoma,Arial,sans-serif; }
}
@page cover { @bottom-center { content: none; } }
body{font:11.5px/1.55 "Segoe UI",Tahoma,Arial,sans-serif;color:#16233a}
.cover{page:cover;page-break-after:always;height:267mm;display:flex;flex-direction:column;justify-content:center;text-align:center;background:#0e2a47;color:#fff;border-radius:10px;padding:30mm 18mm}
.cover .k{font-size:12px;letter-spacing:3px;text-transform:uppercase;color:#8fc1ff;margin:0 0 10px}
.cover h1{font-size:40px;margin:0 0 6px;color:#fff;border:0;padding:0}
.cover h2{font-size:17px;color:#cfe0f2;border:0;padding:0;margin:0 0 4px;page-break-before:auto;font-weight:600}
.cover .ar{direction:rtl;font-size:16px;color:#e0b050;font-weight:700;margin:10px 0 0}
.cover .stats{margin:16px 0 0;font-size:13px;color:#cfe0f2}
.cover .rule{width:60mm;height:2px;background:#e0b050;margin:14px auto}
.toc{page-break-after:always}
.toc h2{page-break-before:auto}
a.tl{display:block;text-decoration:none;color:#16233a;font-size:12.5px;margin:5px 0}
a.tl::after{content:leader('.') target-counter(attr(href), page);color:#5c6b84}
a.tl.sub{margin-inline-start:14px;font-size:11.5px;color:#33415a}
a.tl b{color:#0f4c81}
h1{font-size:22px;color:#0e2a47;margin:0 0 2px}
.sub{color:#5c6b84;font-size:12px;margin:0 0 14px}
h2{font-size:17px;color:#0f4c81;border-bottom:2px solid #dfe6f0;padding-bottom:5px;margin:0 0 10px;page-break-before:always;page-break-after:avoid}
h2.first{page-break-before:auto}
h3{font-size:13px;color:#b8860b;text-transform:uppercase;letter-spacing:.4px;margin:14px 0 6px;page-break-after:avoid}
.fe{border:1px solid #dfe6f0;border-inline-start:4px solid #1663c7;border-radius:8px;padding:7px 10px;margin:7px 0;page-break-inside:avoid;background:#f7f9fc}
.nm{font-weight:800;color:#0f4c81;font-size:12.5px}
.an{color:#b8860b;font-weight:700;font-size:12px}
.f{font-family:Georgia,"Times New Roman",serif;font-size:13px;margin:4px 0;background:#eef3f9;border-radius:6px;padding:5px 9px}
.fr{display:inline-block;vertical-align:middle;text-align:center;margin:0 3px}
.fr .top{display:block;border-bottom:1.2px solid #16233a;padding:0 4px 1px}
.fr .bot{display:block;padding:1px 4px 0}
.vs{font-size:10.5px;color:#5c6b84;margin:2px 0 0}
.wt{font-size:10.5px;color:#33415a;margin:2px 0 0}
.ar{direction:rtl;text-align:right}
table{border-collapse:collapse;width:100%;font-size:10.5px;margin:6px 0}
th,td{border:1px solid #dfe6f0;padding:4px 7px;text-align:start;vertical-align:top}
th{background:#eef3f9;color:#0f4c81;font-size:10px;text-transform:uppercase}
tr{page-break-inside:avoid}
thead{display:table-header-group}
.ge{page-break-inside:avoid;margin:6px 0;border:1px solid #dfe6f0;border-radius:8px;padding:6px 10px}
.ge .t{font-weight:800;color:#0f4c81;font-size:12.5px}
.ge .bd,.bd2{display:inline-block;background:#e8f0fb;color:#1663c7;border-radius:12px;padding:0 8px;font-size:10px;font-weight:700}
.bd2{background:#0f4c81;color:#fff;margin-inline-end:4px}
.ge p{margin:2px 0;font-size:11px}
"""

def fe_html(f, num):
    ch = CH[f['ch']]
    h = '<div class="fe"><span class="bd2">F%02d</span> <span class="nm">%s</span> <span class="an">· %s</span> <span style="float:right;font-size:10px;color:#5c6b84">Chapter %d</span>' % (
        num, escape(f['name']['en']), escape(f['name']['ar']), ch['num'])
    h += '<div class="f">%s</div>' % f['disp']
    if f.get('vars'):
        h += '<div class="vs">' + ' · '.join('<b>%s</b> = %s / %s' % (v['s'], escape(v['en']), escape(v['ar'])) for v in f['vars']) + '</div>'
    h += '<div class="wt"><b>When:</b> %s — <b>Tells:</b> %s</div>' % (escape(f['when']['en']), escape(f['tells']['en']))
    h += '<div class="wt ar"><b>متى:</b> %s — <b>تخبرك:</b> %s</div>' % (escape(f['when']['ar']), escape(f['tells']['ar']))
    return h + '</div>'

parts = ['<html><head><meta charset="utf-8"><style>%s</style></head><body>' % CSS]
parts.append('<div class="cover"><p class="k">FIN211 · Printable Reference</p><h1>Corporate Finance</h1><h2>Formulas · Abbreviations · Definitions</h2><div class="rule"></div><p class="ar">المالية الإدارية — المعادلات والاختصارات والتعريفات</p><p class="stats">%d Formulas · %d Abbreviations · %d Definitions — bilingual EN/AR<br>%d معادلة · %d اختصارًا · %d تعريفًا — باللغتين</p><p class="stats" style="margin-top:26px;font-size:11px;opacity:.8">Built from the course lecture materials (Lectures 1–13)<br>مبني من محاضرات المقرر (١–١٣)</p><p class="stats" style="margin-top:12px;color:#e0b050;font-weight:700">Prepared by 领袖 - EBRA · إعداد</p></div>' % (
    len(DATA['formulas']), len(DATA['abbrev']), len(DATA['glossary']) + len(SUPPL),
    len(DATA['formulas']), len(DATA['abbrev']), len(DATA['glossary']) + len(SUPPL)))
parts.append('<div class="toc"><h2 class="first">Contents — الفهرس</h2>')
parts.append('<a class="tl" href="#secA"><b>A · Formula Bank — بنك المعادلات</b></a>')
for c in DATA['chapters']:
    if any(f['ch'] == c['id'] for f in DATA['formulas']):
        parts.append('<a class="tl sub" href="#ch%s">%d · %s · %s</a>' % (c['id'], c['num'], escape(c['title']['en']), escape(c['title']['ar'])))
parts.append('<a class="tl" href="#secB"><b>B · Abbreviations — الاختصارات</b></a>')
parts.append('<a class="tl" href="#secC"><b>C · Definitions — التعريفات</b></a>')
parts.append('</div>')
parts.append('<p class="sub">Complete printable reference: every formula (with variables, when to use and what it tells you), every abbreviation, and every definition from the course lectures — bilingual EN/AR. · مرجع طباعي كامل: كل المعادلات والاختصارات والتعريفات من محاضرات المقرر — باللغتين.</p>')

parts.append('<h2 class="first" id="secA">A · Formula Bank — بنك المعادلات</h2>')
for c in DATA['chapters']:
    fs = [f for f in DATA['formulas'] if f['ch'] == c['id']]
    if not fs: continue
    parts.append('<h3 id="ch%s">Chapter %d — %s · %s</h3>' % (c['id'], c['num'], escape(c['title']['en']), escape(c['title']['ar'])))
    for f in fs:
        FI[f['id']] = len(FI) + 1
        parts.append(fe_html(f, FI[f['id']]))

parts.append('<h2 id="secB">B · Abbreviations — الاختصارات</h2>')
parts.append('<table><thead><tr><th style="width:8%">#</th><th style="width:13%">Abbr.</th><th style="width:34%">Full name (English)</th><th style="width:33%">بالعربية</th><th style="width:12%">Lesson</th></tr></thead><tbody>')
for a in DATA['abbrev']:
    AI[a['a']] = len(AI) + 1
    parts.append('<tr><td>A%02d</td><td><b>%s</b></td><td>%s</td><td class="ar">%s</td><td>%d</td></tr>' % (
        AI[a['a']], escape(a['a']), escape(a['en']), escape(a['ar']), int(a['l'][1:])))
parts.append('</tbody></table>')

parts.append('<h2 id="secC">C · Definitions — التعريفات</h2>')
def def_html(num, term_en, abbr, def_en, term_ar, def_ar):
    ab = '' if abbr in ('—', '') else ' <span class="bd">%s</span>' % escape(abbr)
    return '<div class="ge"><span class="bd2">D%02d</span> <span class="t">%s</span>%s<p>%s</p><p class="ar"><b>%s:</b> %s</p></div>' % (
        num, escape(term_en), ab, escape(def_en), escape(term_ar), escape(def_ar))
for g in DATA['glossary']:
    DI[g['t']['en']] = len(DI) + 1
    parts.append(def_html(DI[g['t']['en']], g['t']['en'], g['ab'], g['d']['en'], g['t']['ar'], g['d']['ar']))
for sp in SUPPL:
    DI[sp[1]] = len(DI) + 1
    parts.append(def_html(DI[sp[1]], sp[1], sp[0], sp[2], sp[3], sp[4]))
parts.append('</body></html>')
html = '\n'.join(parts)
io.open(os.path.join(HERE, '_ref.html'), 'w', encoding='utf-8').write(html)

from weasyprint import HTML
HTML(string=html).write_pdf(OUT_PDF)
print('PDF written:', OUT_PDF)

# ============================== DOCX ==============================
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(14)
sec.top_margin = sec.bottom_margin = Mm(13)
cov = doc.add_heading('Corporate Finance', 0)
cov.alignment = 1
ps = doc.add_paragraph('FIN211 · Formulas · Abbreviations · Definitions')
ps.alignment = 1; ps.runs[0].bold = True
ps2 = doc.add_paragraph('المالية الإدارية — المعادلات والاختصارات والتعريفات')
ps2.alignment = 1
ps3 = doc.add_paragraph('%d Formulas · %d Abbreviations · %d Definitions — bilingual EN/AR' % (
    len(DATA['formulas']), len(DATA['abbrev']), len(DATA['glossary']) + len(SUPPL)))
ps3.alignment = 1; ps3.runs[0].italic = True
pc = doc.add_paragraph('Prepared by 领袖 - EBRA · إعداد'); pc.alignment = 1; pc.runs[0].bold = True
doc.add_page_break()
def _field(par, code):
    r = par.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = code
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    r._r.append(f1); r._r.append(it); r._r.append(f2)
fp = sec.footer.paragraphs[0]; fp.alignment = 1
_field(fp, 'PAGE'); fp.add_run(' / '); _field(fp, 'NUMPAGES')

def ar_par(text, bold_prefix=None):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr(); bidi = OxmlElement('w:bidi'); pPr.append(bidi)
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(2)
    return p

h = doc.add_heading('A · Formula Bank — بنك المعادلات', 1)
first_ch = True
for c in DATA['chapters']:
    fs = [f for f in DATA['formulas'] if f['ch'] == c['id']]
    if not fs: continue
    hc = doc.add_heading('Chapter %d — %s · %s' % (c['num'], c['title']['en'], c['title']['ar']), 2)
    hc.paragraph_format.keep_with_next = True
    if not first_ch: hc.paragraph_format.page_break_before = True
    first_ch = False
    for f in fs:
        pn = doc.add_paragraph(); pn.paragraph_format.keep_with_next = True
        r0 = pn.add_run('F%02d  ' % FI[f['id']]); r0.bold = True; r0.font.color.rgb = RGBColor(0x0f, 0x4c, 0x81)
        r = pn.add_run(f['name']['en']); r.bold = True; r.font.color.rgb = RGBColor(0x0f, 0x4c, 0x81)
        r2 = pn.add_run('  ·  ' + f['name']['ar']); r2.bold = True; r2.font.color.rgb = RGBColor(0xb8, 0x86, 0x0b)
        pf = doc.add_paragraph(lin(f['disp']))
        pf.paragraph_format.left_indent = Mm(5); pf.paragraph_format.space_after = Pt(2)
        for run in pf.runs: run.font.name = 'Cambria'
        if f.get('vars'):
            pv = doc.add_paragraph('; '.join('%s = %s / %s' % (re.sub(r'<[^>]+>','',v['s']), v['en'], v['ar']) for v in f['vars']))
            pv.runs[0].font.size = Pt(9); pv.runs[0].font.color.rgb = RGBColor(0x5c, 0x6b, 0x84)
            pv.paragraph_format.space_after = Pt(1)
        pw = doc.add_paragraph('When: %s — Tells: %s' % (f['when']['en'], f['tells']['en']))
        pw.runs[0].font.size = Pt(9); pw.paragraph_format.space_after = Pt(1)
        ar_par('متى: %s — تخبرك: %s' % (f['when']['ar'], f['tells']['ar']))
        doc.add_paragraph('').paragraph_format.space_after = Pt(2)

doc.add_heading('B · Abbreviations — الاختصارات', 1).paragraph_format.page_break_before = True
tb = doc.add_table(rows=1, cols=5); tb.style = 'Table Grid'
hdr = tb.rows[0]; hdr.cant_split = True
trPr = hdr._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); trPr.append(th)
for i, ttxt in enumerate(['#', 'Abbr.', 'Full name (English)', 'بالعربية', 'Lesson']):
    hdr.cells[i].text = ttxt
    for rr in hdr.cells[i].paragraphs[0].runs: rr.bold = True
for a in DATA['abbrev']:
    ln = [l for l in DATA['lessons'] if l['id'] == a['l']][0]
    row = tb.add_row(); row.cant_split = True
    AI[a['a']] = AI.get(a['a'], len(AI) + 1)
    row.cells[0].text = 'A%02d' % AI[a['a']]
    row.cells[1].text = a['a']; row.cells[2].text = a['en']; row.cells[3].text = a['ar']
    row.cells[4].text = ln['id'].replace('l', '')
    for pr in row.cells[3].paragraphs:
        pPr = pr._p.get_or_add_pPr(); bidi = OxmlElement('w:bidi'); pPr.append(bidi)

doc.add_heading('C · Definitions — التعريفات', 1).paragraph_format.page_break_before = True
for g in DATA['glossary']:
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(1)
    DI[g['t']['en']] = DI.get(g['t']['en'], len(DI) + 1)
    r0 = p.add_run('D%02d  ' % DI[g['t']['en']]); r0.bold = True
    r = p.add_run(g['t']['en']); r.bold = True; r.font.color.rgb = RGBColor(0x0f, 0x4c, 0x81)
    if g['ab'] != '—':
        r2 = p.add_run('  (' + g['ab'] + ')'); r2.bold = True; r2.font.size = Pt(9)
    pe = doc.add_paragraph(g['d']['en']); pe.paragraph_format.space_after = Pt(1)
    ar_par(g['d']['ar'], bold_prefix=g['t']['ar'] + ': ')
for sp in SUPPL:
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(1)
    DI[sp[1]] = DI.get(sp[1], len(DI) + 1)
    r0 = p.add_run('D%02d  ' % DI[sp[1]]); r0.bold = True
    r = p.add_run(sp[1]); r.bold = True; r.font.color.rgb = RGBColor(0x0f, 0x4c, 0x81)
    r2 = p.add_run('  (' + sp[0] + ')'); r2.bold = True; r2.font.size = Pt(9)
    pe = doc.add_paragraph(sp[2]); pe.paragraph_format.space_after = Pt(1)
    ar_par(sp[4], bold_prefix=sp[3] + ': ')
doc.save(OUT_DOCX)
print('DOCX written:', OUT_DOCX)
